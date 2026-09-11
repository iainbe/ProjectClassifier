#!/usr/bin/env python3
"""Extract text from G2 Phase 1 pre-2020 cohort source documents."""
import os
from docx import Document

DRIVE = os.path.expanduser("~/Library/CloudStorage/GoogleDrive-iain@thefifthsector.co.uk/My Drive")
OUT = os.path.join(DRIVE, "Website 2026/spillover-toolkit/extracted_text")

# Key deliverables to extract for full reading
sources = [
    # Kirklees 2015 — BOP report (may be in zip, but extract what we have)
    ("SRC-G2-001", "Archive projects/2015 Projects/2015 Kirklees/150416 Kirklees Results.xlsx", "xlsx"),
    # ELFC — final report (Iain was lead)
    ("SRC-G2-002", "Archive projects/2017 Projects/2017 East London Fashion Cluster/170302_ELFC_Report_DIGITAL.docx", "docx"),
    # ELFC — summary report
    ("SRC-G2-003", "Archive projects/2017 Projects/2017 East London Fashion Cluster/170314_ELFC_SummaryReport_PRESSQUALITY_FINAL.docx", "docx"),
    # ELFC — case for intervention (Iain authored)
    ("SRC-G2-004", "Archive projects/2017 Projects/2017 East London Fashion Cluster/161005 East London Fashion Cluster case for intervention v2_IB.docx", "docx"),
    # ELFC — action plan
    ("SRC-G2-005", "Archive projects/2017 Projects/2017 East London Fashion Cluster/170314 ELFC Action Plan final v3.0 clean.docx", "docx"),
]

for source_id, rel_path, ftype in sources:
    full_path = os.path.join(DRIVE, rel_path)
    out_file = os.path.join(OUT, f"{source_id}_v1_full.txt")

    if os.path.exists(out_file):
        print(f"SKIP {source_id}: output already exists")
        continue

    if not os.path.exists(full_path):
        print(f"MISSING {source_id}: {full_path}")
        continue

    try:
        if ftype == "docx":
            doc = Document(full_path)
            lines = []
            para_count = 0
            table_count = 0

            for para in doc.paragraphs:
                lines.append(para.text)
                para_count += 1

            for table in doc.tables:
                table_count += 1
                lines.append(f"\n[TABLE {table_count}]")
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append(" | ".join(cells))
                lines.append("[/TABLE]\n")

            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)

            print(f"OK {source_id}: {para_count} paras, {table_count} tables, {len(content)} chars")
        elif ftype == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(full_path, data_only=True)
            lines = []
            sheet_count = 0

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_count += 1
                lines.append(f"\n[SHEET {sheet_count}: {sheet_name}]")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    lines.append(" | ".join(cells))
                lines.append(f"[/SHEET {sheet_count}]\n")

            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)

            print(f"OK {source_id}: {sheet_count} sheets, {len(content)} chars")
        elif ftype == "pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(full_path)
            lines = []
            for i, page in enumerate(reader.pages):
                lines.append(f"\n[PAGE {i+1}]")
                text = page.extract_text()
                if text:
                    lines.append(text)
            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"OK {source_id}: {len(reader.pages)} pages, {len(content)} chars")

        print(f"  -> {out_file}")

    except Exception as e:
        print(f"ERROR {source_id}: {e}")
        import traceback
        traceback.print_exc()

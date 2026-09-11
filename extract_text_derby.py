#!/usr/bin/env python3
"""Extract text from T05-DERBY sources."""
import os
from docx import Document

DRIVE = os.path.expanduser("~/Library/CloudStorage/GoogleDrive-iain@thefifthsector.co.uk/My Drive")
OUT = os.path.join(DRIVE, "Website 2026/spillover-toolkit/extracted_text")

sources = [
    ("SRC-R3-04", "Archive proposals/2026 Proposals lost/2026 Derby/ENQ1304 - Quick Quote Request for Quotation (RFQ).docx", "docx"),
    ("SRC-R3-05", "Archive proposals/2026 Proposals lost/2026 Derby/Derby_ENQ1304_Response_TheFifthSector_v3.docx", "docx"),
    ("SRC-R3-06", "Archive proposals/2026 Proposals lost/2026 Derby/260624 Hannah Barker Derby.txt", "txt"),
    ("SRC-R3-07", "Archive proposals/2026 Proposals lost/2026 Derby/Cultural Data Strategy Meeting_otter_ai_transcript.txt", "txt"),
]

for source_id, rel_path, ftype in sources:
    full_path = os.path.join(DRIVE, rel_path)
    out_file = os.path.join(OUT, f"{source_id}_v1_full.txt")

    if os.path.exists(out_file):
        print(f"SKIP {source_id}: output already exists")
        continue

    if not os.path.exists(full_path):
        print(f"MISSING {source_id}: {full_path}")
        continue

    try:
        if ftype == "docx":
            doc = Document(full_path)
            lines = []
            para_count = 0
            table_count = 0

            for para in doc.paragraphs:
                lines.append(para.text)
                para_count += 1

            for table in doc.tables:
                table_count += 1
                lines.append(f"\n[TABLE {table_count}]")
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append(" | ".join(cells))
                lines.append("[/TABLE]\n")

            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)

            print(f"OK {source_id}: {para_count} paras, {table_count} tables, {len(content)} chars")
        elif ftype == "txt":
            import shutil
            shutil.copy2(full_path, out_file)
            with open(out_file, "r", encoding="utf-8") as f:
                content = f.read()
            print(f"OK {source_id}: {len(content)} chars (copied txt)")

        print(f"  -> {out_file}")

    except Exception as e:
        print(f"ERROR {source_id}: {e}")

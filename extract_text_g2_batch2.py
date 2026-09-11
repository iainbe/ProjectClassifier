#!/usr/bin/env python3
"""Extract text from G2 Phase 1 pre-2020 remaining sources (lighter touch)."""
import os
from docx import Document

DRIVE = os.path.expanduser("~/Library/CloudStorage/GoogleDrive-iain@thefifthsector.co.uk/My Drive")
OUT = os.path.join(DRIVE, "Website 2026/spillover-toolkit/extracted_text")

sources = [
    # Creativeworks London — final report
    ("SRC-G2-006", "Archive projects/2016 Projects/2016 Creativeworks London/160601_CWL_Final.docx", "docx"),
    # Digital Catapult Immersive — final report (the IB-marked version)
    ("SRC-G2-007", "Archive projects/2018 Projects/2018 Digital Catapult immersive/180318_Digital_Catapault_Report_IB.pdf", "pdf"),
    # Digital Catapult — technical report
    ("SRC-G2-008", "Archive projects/2018 Projects/2018 Digital Catapult immersive/BEI-WP7_Technical-report.docx", "docx"),
    # Liverpool City of Music — strategy draft (the most complete version)
    ("SRC-G2-009", "Archive projects/2017 Projects/2017 Liverpool City of Music/180109 Liverpool City of Music strategy draft final report.docx", "docx"),
    # Liverpool City of Music — final strategy
    ("SRC-G2-010", "Archive projects/2017 Projects/2017 Liverpool City of Music/Liverpool City of Music strategy February 2018.docx", "docx"),
]

for source_id, rel_path, ftype in sources:
    full_path = os.path.join(DRIVE, rel_path)
    out_file = os.path.join(OUT, f"{source_id}_v1_full.txt")

    if os.path.exists(out_file):
        print(f"SKIP {source_id}: output already exists")
        continue

    if not os.path.exists(full_path):
        print(f"MISSING {source_id}: {full_path}")
        continue

    try:
        if ftype == "docx":
            doc = Document(full_path)
            lines = []
            para_count = 0
            table_count = 0

            for para in doc.paragraphs:
                lines.append(para.text)
                para_count += 1

            for table in doc.tables:
                table_count += 1
                lines.append(f"\n[TABLE {table_count}]")
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append(" | ".join(cells))
                lines.append("[/TABLE]\n")

            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)

            print(f"OK {source_id}: {para_count} paras, {table_count} tables, {len(content)} chars")
        elif ftype == "pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(full_path)
            lines = []
            for i, page in enumerate(reader.pages):
                lines.append(f"\n[PAGE {i+1}]")
                text = page.extract_text()
                if text:
                    lines.append(text)
            content = "\n".join(lines)
            with open(out_file, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"OK {source_id}: {len(reader.pages)} pages, {len(content)} chars")

        print(f"  -> {out_file}")

    except Exception as e:
        print(f"ERROR {source_id}: {e}")
        import traceback
        traceback.print_exc()
